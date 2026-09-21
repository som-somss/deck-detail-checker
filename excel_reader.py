from openpyxl import load_workbook
from models import Spec

def num(v):
    try:
        if v in (None,""):return None
        return float(v)
    except:return None

def read_excel_specs(path):
    wb=load_workbook(path,data_only=True,read_only=True)
    if "제원입력" not in wb.sheetnames:
        raise ValueError("Excel에서 '제원입력' 시트를 찾지 못했습니다.")
    ws=wb["제원입력"]; out={}
    for r in range(1,ws.max_row+1):
        typ=ws.cell(r,2).value
        if not isinstance(typ,str) or not typ.strip().upper().startswith("M"):continue
        typ=typ.strip()
        c,e,g,i=map(num,(ws.cell(r,3).value,ws.cell(r,5).value,ws.cell(r,7).value,ws.cell(r,9).value))
        if c is not None and e is not None:b1,b2=c,e
        elif e is not None and i is not None:b1,b2=e,i
        else:
            b1=e if e is not None else c
            b2=i if i is not None else b1
        if b1 is not None and g is not None:out[typ]=Spec(typ,b1,b2,g)
    return out
